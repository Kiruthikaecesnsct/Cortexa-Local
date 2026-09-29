import io

import docx

from ingestion.application.text_normalizer import normalize_text
from ingestion.domain.errors.parser_errors import CorruptedFileError
from ingestion.domain.models.parsed_document import ParsedDocument


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs]
    except Exception as e:
        raise CorruptedFileError(f"DOCX parsing failed: {e}") from e

    text = normalize_text("\n".join(paragraphs))
    return ParsedDocument(text=text, pages=[])
