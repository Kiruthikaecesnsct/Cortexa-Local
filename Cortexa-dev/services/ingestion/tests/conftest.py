import io

import docx
import pytest

MINIMAL_PDF = b"""%PDF-1.4
1 0 obj<</Type /Catalog /Pages 2 0 R>>endobj
2 0 obj<</Type /Pages /Kids [3 0 R] /Count 1>>endobj
3 0 obj<</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]
/Contents 4 0 R /Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 44>>
stream
BT /F1 12 Tf 100 700 Td (Hello  World) Tj ET
endstream
endobj
5 0 obj<</Type /Font /Subtype /Type1 /BaseFont /Helvetica>>endobj
xref
0 6
0000000000 65535 f\r
0000000009 00000 n\r
0000000058 00000 n\r
0000000115 00000 n\r
0000000266 00000 n\r
0000000360 00000 n\r
trailer<</Size 6 /Root 1 0 R>>
startxref
441
%%EOF"""


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    return MINIMAL_PDF


@pytest.fixture
def sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def corrupted_bytes() -> bytes:
    return b"this is not a valid file format at all"


def make_multi_page_pdf(pages_text: list[str]) -> bytes:
    catalog = b"<</Type /Catalog /Pages 2 0 R>>"
    kids = " ".join(f"{3 + i} 0 R" for i in range(len(pages_text)))
    pages_obj = f"<</Type /Pages /Kids [{kids}] /Count {len(pages_text)}>>".encode()
    content_obj_start = 3 + len(pages_text)
    font_num = content_obj_start + len(pages_text)

    page_objs = []
    content_objs = []
    for i, text in enumerate(pages_text):
        content_num = content_obj_start + i
        page_objs.append(
            f"<</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_num} 0 R /Resources<</Font<</F1 {font_num} 0 R>>>>>>".encode()
        )
        stream = f"BT /F1 12 Tf 100 700 Td ({text}) Tj ET".encode()
        content_objs.append(
            b"<</Length " + str(len(stream)).encode() + b">>\nstream\n" + stream + b"\nendstream\n"
        )
    font_obj = b"<</Type /Font /Subtype /Type1 /BaseFont /Helvetica>>"
    all_objs = [catalog, pages_obj, *page_objs, *content_objs, font_obj]

    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n")
    offsets = [0]
    for idx, obj in enumerate(all_objs, start=1):
        offsets.append(buf.tell())
        buf.write(f"{idx} 0 obj".encode() + obj + b"\nendobj\n")
    xref_offset = buf.tell()
    n = len(all_objs) + 1
    buf.write(f"xref\n0 {n}\n".encode())
    buf.write(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        buf.write(f"{off:010d} 00000 n \n".encode())
    buf.write(f"trailer<</Size {n} /Root 1 0 R>>\nstartxref\n{xref_offset}\n%%EOF".encode())
    return buf.getvalue()


@pytest.fixture
def two_page_pdf_bytes() -> bytes:
    return make_multi_page_pdf(["Page one text here", "Page two content words"])
